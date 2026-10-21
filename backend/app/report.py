"""判定結果一覧（マスキング対象一覧・要確認一覧）の Excel 出力。

元の値は出力せず、マスク後の例のみを載せる（一覧自体を共有・保管しても安全なように）。
"""

from __future__ import annotations

import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .findings import effective_rule, masked_examples
from .masking import RULE_DEFINITIONS, display_width
from .schemas import Decision, Finding

_RULE_LABELS = {r["type"]: r["label"] for r in RULE_DEFINITIONS}
FONT = "游ゴシック"

COLUMNS = [
    ("No", 6),
    ("ファイル・表", 28),
    ("場所", 26),
    ("対象", 22),
    ("論理名", 18),
    ("データ型", 16),
    ("種類", 22),
    ("確からしさ", 10),
    ("判断", 14),
    ("処理（マスキング規則）", 26),
    ("件数", 8),
    ("マスク後の例", 34),
    ("判定理由", 60),
]


def _row(i: int, f: Finding, decisions: dict[str, Decision]) -> list:
    rule = effective_rule(f, decisions)
    d = decisions.get(f.id)
    if f.level == "auto":
        changed = d is not None and (not d.mask or (d.rule is not None and d.rule != f.rule))
        judged = "自動（人が変更）" if changed else "自動"
    else:
        judged = "人が判断" if d is not None else "未判断"
    action = _RULE_LABELS.get(rule.type, rule.type) if rule else "マスクしない"
    target = f.column if f.target == "column" else (f.text or "文中の語句")
    return [
        i,
        f.file,
        f.location or "",
        ("列 " if f.target == "column" else "文中 ") + (target or ""),
        f.column_comment or "",
        f.data_type or "",
        f.category_label,
        round(f.score, 2),
        judged,
        action,
        f.occurrences,
        " / ".join(masked_examples(rule, f.samples)) if f.samples else "",
        "、".join(f.reasons),
    ]


def build_report_xlsx(title: str, findings: list[Finding], decisions: dict[str, Decision], source_label: str | None) -> bytes:
    wb = Workbook()
    font = Font(name=FONT, size=10)
    head_font = Font(name=FONT, size=10, bold=True)
    fill = PatternFill("solid", fgColor="DDEBF7")
    fill_keep = PatternFill("solid", fgColor="EDEDED")
    thin = Side(style="thin", color="A6A6A6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(vertical="top", wrap_text=True)

    summary = wb.active
    summary.title = "概要"
    auto = [f for f in findings if f.level == "auto"]
    review = [f for f in findings if f.level == "review"]
    undecided = sum(1 for f in review if f.id not in decisions)
    rows = [
        [title],
        [],
        ["出力日時", dt.datetime.now().strftime("%Y/%m/%d %H:%M:%S")],
        ["対象", source_label or ""],
        ["マスキング対象（自動判定）", len(auto)],
        ["要確認（人が判断）", len(review)],
        ["　うち未判断", undecided],
        ["マスクしない", sum(1 for f in findings if effective_rule(f, decisions) is None)],
        [],
        ["※ 元の値は記載していません（マスク後の例のみ）。"],
        ["※ 判定は列名・論理名・値の形式による推定です。最終的な判断は人が行ってください。"],
    ]
    for r in rows:
        summary.append(r)
    for row in summary.iter_rows():
        for c in row:
            c.font = font
    summary["A1"].font = Font(name=FONT, size=14, bold=True)
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 50

    for sheet_title, items in (("マスキング対象一覧", auto), ("要確認一覧", review)):
        ws = wb.create_sheet(sheet_title)
        ws.append([c for c, _ in COLUMNS])
        for i, f in enumerate(items, start=1):
            ws.append(_row(i, f, decisions))
        for idx, (_, width) in enumerate(COLUMNS, start=1):
            letter = get_column_letter(idx)
            longest = max((display_width(str(ws.cell(r, idx).value or "")) for r in range(2, min(ws.max_row, 200) + 1)), default=0)
            ws.column_dimensions[letter].width = min(max(width, longest * 0.9 + 2), 80 if idx == len(COLUMNS) else 50)
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for c in row:
                c.font = font
                c.border = border
                c.alignment = wrap
        for c in ws[1]:
            c.font = head_font
            c.fill = fill
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 10).value == "マスクしない":
                for c in ws[r]:
                    c.fill = fill_keep
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(ws.max_row, 1)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
