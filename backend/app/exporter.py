"""エクスポート処理（Excel / マスク済み INSERT 文 / マスキング用 UPDATE 文）。

- openpyxl の write_only モードでストリーミング出力し、大量件数でもメモリを圧迫しない。
- 1 シートの上限（1,048,576 行）を超える場合は「表名_2」のように自動でシートを分割する。
- 日本語フォント・全角文字幅を考慮した列幅設定、見出し行の固定、オートフィルタを付与する。
- 処理はバックグラウンドスレッドで実行し、進捗をポーリングで取得できる。
"""

from __future__ import annotations

import datetime as dt
import json
import logging
import re
import threading
import time
import uuid
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.indexed_list import IndexedList

from . import sqlgen, store
from .config import export_dir, load_settings
from .db import DbError
from .masking import RULE_DEFINITIONS, build_masker, display_width
from .schemas import (
    ColumnInfo,
    ExportJobStatus,
    ExportOptions,
    ExportRequest,
    TableExportSpec,
    TableProgress,
)

log = logging.getLogger(__name__)

EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_CELL_CHARS = 32_767
# Excel の数値精度（有効桁 15 桁）を超える数値は文字列で出力する
EXCEL_MAX_DIGITS = 15
FETCH_SIZE = 5000

_RULE_LABELS = {r["type"]: r["label"] for r in RULE_DEFINITIONS}


# ---------------------------------------------------------------------------
# 値の変換
# ---------------------------------------------------------------------------


def to_excel_value(value: Any) -> Any:
    """DB から取得した値を Excel に書き込める形に変換する。"""
    if value is None:
        return None
    if isinstance(value, str):
        text = ILLEGAL_CHARACTERS_RE.sub("", value)
        return text[:EXCEL_MAX_CELL_CHARS]
    if isinstance(value, bool):
        return value
    if isinstance(value, Decimal):
        if not value.is_finite():
            return str(value)
        digits = len(value.as_tuple().digits)
        if digits > EXCEL_MAX_DIGITS:
            # 桁数の多い ID・口座番号等は精度が落ちないよう文字列で出力
            return format(value, "f")
        return int(value) if value == value.to_integral_value() else float(value)
    if isinstance(value, int):
        return value if abs(value) < 10**EXCEL_MAX_DIGITS else str(value)
    if isinstance(value, float):
        return value
    if isinstance(value, dt.datetime):
        # Excel はタイムゾーン付き日時を扱えないため、PC のローカル時刻（日本なら JST）に
        # 変換してからタイムゾーン情報を除去する
        return value.astimezone().replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, (dt.date, dt.time)):
        return value
    if isinstance(value, dt.timedelta):
        return str(value)
    if isinstance(value, (bytes, bytearray, memoryview)):
        data = bytes(value)
        if len(data) <= 1000:
            return data.hex().upper()
        return f"(バイナリ {len(data):,} バイト)"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)[:EXCEL_MAX_CELL_CHARS]
    return str(value)[:EXCEL_MAX_CELL_CHARS]


def to_json_value(value: Any) -> Any:
    """プレビュー表示（JSON）用に変換する。"""
    value = to_excel_value(value)
    if isinstance(value, dt.datetime):
        return value.strftime("%Y/%m/%d %H:%M:%S")
    if isinstance(value, dt.date):
        return value.strftime("%Y/%m/%d")
    if isinstance(value, dt.time):
        return value.strftime("%H:%M:%S")
    return value


# ---------------------------------------------------------------------------
# 列の準備（メタデータ照合・見出し・マスク関数）
# ---------------------------------------------------------------------------


@dataclass
class PreparedColumn:
    name: str
    logical: str
    data_type: str
    masker: Callable[[Any], Any] | None
    mask_label: str | None


def prepare_columns(spec: TableExportSpec, meta: list[ColumnInfo], options: ExportOptions) -> list[PreparedColumn]:
    """指定列が実在するか確認し、見出し・マスク関数を準備する。

    SQL に埋め込む列名は必ずここで DB のメタデータと照合済みのものに限定する。
    """
    by_name = {c.name: c for c in meta}
    prepared: list[PreparedColumn] = []
    for col in spec.columns:
        info = by_name.get(col.name)
        if info is None:
            raise DbError(f"列 {col.name} は {spec.schema_name}.{spec.table} に存在しません。")
        try:
            masker = build_masker(col.mask, options.hash_salt)
        except ValueError as exc:
            raise DbError(str(exc)) from exc
        logical = col.header or info.comment or info.name
        mask_label = None
        if masker is not None and col.mask is not None:
            mask_label = _RULE_LABELS.get(col.mask.type, col.mask.type)
        prepared.append(PreparedColumn(info.name, logical, info.data_type, masker, mask_label))
    return prepared


def header_rows(columns: list[PreparedColumn], mode: str) -> list[list[str]]:
    if mode == "physical":
        return [[c.name for c in columns]]
    if mode == "both":
        return [[c.logical for c in columns], [c.name for c in columns]]
    return [[c.logical for c in columns]]


def transform_row(row: tuple, columns: list[PreparedColumn]) -> list[Any]:
    out = []
    for value, col in zip(row, columns):
        if col.masker is not None:
            value = col.masker(value)
        out.append(to_excel_value(value))
    return out


# ---------------------------------------------------------------------------
# Excel 書き込み
# ---------------------------------------------------------------------------

_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_INVALID_FILE_CHARS = re.compile(r'[\\/:*?"<>|\x00-\x1f]')


def safe_sheet_name(name: str, used: set[str]) -> str:
    """Excel のシート名制約（31文字以内・禁止文字・重複不可）に合わせる。"""
    base = _INVALID_SHEET_CHARS.sub("_", name).strip("'") or "Sheet"
    candidate = base[:31]
    n = 2
    while candidate.lower() in used:
        suffix = f"_{n}"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


FORMAT_EXT = {"xlsx": ".xlsx", "sql_insert": ".sql", "sql_update": ".sql"}
MEDIA_TYPES = {
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".sql": "application/sql",
}


def safe_file_name(name: str, ext: str = ".xlsx") -> str:
    name = _INVALID_FILE_CHARS.sub("_", name).strip(" .") or "export"
    if not name.lower().endswith(ext):
        name += ext
    return name


class ExcelBook:
    """write_only モードの Workbook を日本語向け書式でラップする。"""

    def __init__(self, options: ExportOptions):
        self.options = options
        self.wb = Workbook(write_only=True)
        self._apply_default_font()
        self.used_names: set[str] = set()
        self.font_header = Font(name=options.font_name, size=options.font_size, bold=True)
        self.fill_header = PatternFill("solid", fgColor="DDEBF7")
        self.fill_masked = PatternFill("solid", fgColor="FCE4D6")
        thin = Side(style="thin", color="A6A6A6")
        self.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        self.align_header = Alignment(vertical="center", wrap_text=False)

    def _apply_default_font(self) -> None:
        """ブック全体の既定フォント（標準スタイル）を日本語フォントに変更する。

        write_only モードではセル毎のスタイル設定が高コストなため、既定フォントを差し替える。
        """
        font = Font(name=self.options.font_name, size=self.options.font_size, family=3, charset=128)
        self.wb._fonts = IndexedList([font])  # noqa: SLF001
        for style in self.wb._named_styles:  # noqa: SLF001
            if style.name == "Normal":
                style.font = font

    def _header_cell(self, ws: Any, value: str, masked: bool) -> WriteOnlyCell:
        cell = WriteOnlyCell(ws, value=value)
        cell.font = self.font_header
        cell.fill = self.fill_masked if masked else self.fill_header
        cell.border = self.border
        cell.alignment = self.align_header
        return cell

    def new_sheet(self, title: str, columns: list[PreparedColumn], sample: list[list[Any]]) -> Any:
        ws = self.wb.create_sheet(safe_sheet_name(title, self.used_names))
        headers = header_rows(columns, self.options.header_mode)
        # 列幅は「見出し」と「先頭データのサンプル」から全角を考慮して算出する
        for idx, col in enumerate(columns):
            texts = [h[idx] for h in headers] + [_display_text(r[idx]) for r in sample[:200]]
            width = max((display_width(t) for t in texts if t), default=8)
            ws.column_dimensions[get_column_letter(idx + 1)].width = min(max(width * 1.1 + 2, 8), 60)
        ws.freeze_panes = f"A{len(headers) + 1}"
        for row in headers:
            ws.append([self._header_cell(ws, v, columns[i].masker is not None) for i, v in enumerate(row)])
        return ws

    def append(self, ws: Any, values: list[Any]) -> None:
        cells: list[Any] = []
        for v in values:
            if isinstance(v, dt.datetime):
                c = WriteOnlyCell(ws, value=v)
                c.number_format = self.options.datetime_format
                cells.append(c)
            elif isinstance(v, dt.date):
                c = WriteOnlyCell(ws, value=v)
                c.number_format = self.options.date_format
                cells.append(c)
            else:
                cells.append(v)
        ws.append(cells)

    def finish_sheet(self, ws: Any, n_cols: int, n_rows: int) -> None:
        header_count = 2 if self.options.header_mode == "both" else 1
        last_row = header_count + n_rows
        ws.auto_filter.ref = f"A{header_count}:{get_column_letter(n_cols)}{max(last_row, header_count)}"

    def add_info_sheet(self, rows: list[list[Any]]) -> None:
        ws = self.wb.create_sheet(safe_sheet_name("エクスポート情報", self.used_names))
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 40
        ws.column_dimensions["C"].width = 30
        ws.column_dimensions["D"].width = 30
        for r in rows:
            ws.append(r)

    def save(self, path: Path) -> None:
        self.wb.save(path)


def _display_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return "0000/00/00 00:00:00"
    if isinstance(value, dt.date):
        return "0000/00/00"
    return str(value)


# ---------------------------------------------------------------------------
# エクスポートジョブ
# ---------------------------------------------------------------------------


@dataclass
class ExportJob:
    job_id: str
    request: ExportRequest
    file_name: str
    path: Path
    status: str = "pending"
    message: str | None = None
    tables: list[TableProgress] = field(default_factory=list)
    started_at: dt.datetime | None = None
    finished_at: dt.datetime | None = None
    cancel_event: threading.Event = field(default_factory=threading.Event)

    def to_status(self) -> ExportJobStatus:
        return ExportJobStatus(
            job_id=self.job_id,
            status=self.status,  # type: ignore[arg-type]
            message=self.message,
            file_name=self.file_name,
            tables=self.tables,
            started_at=self.started_at.isoformat(timespec="seconds") if self.started_at else None,
            finished_at=self.finished_at.isoformat(timespec="seconds") if self.finished_at else None,
        )


class Cancelled(Exception):
    pass


_jobs: dict[str, ExportJob] = {}
_jobs_lock = threading.Lock()


def get_job(job_id: str) -> ExportJob:
    with _jobs_lock:
        job = _jobs.get(job_id)
    if job is None:
        raise DbError("指定されたエクスポート処理が見つかりません。")
    return job


def cleanup_old_files() -> None:
    """保持期間を過ぎたエクスポートファイルを削除する。"""
    retention = int(load_settings().get("export_retention_seconds") or 0)
    if retention <= 0:
        return
    limit = time.time() - retention
    for p in export_dir().iterdir():
        try:
            if p.stat().st_mtime < limit:
                p.unlink()
        except OSError:
            pass
    with _jobs_lock:
        for jid in [j.job_id for j in _jobs.values() if j.finished_at and not j.path.exists()]:
            _jobs.pop(jid, None)


def start_export(req: ExportRequest) -> ExportJob:
    cleanup_old_files()
    now = dt.datetime.now()
    fmt = req.options.format
    ext = FORMAT_EXT[fmt]
    default = req.tables[0].table if len(req.tables) == 1 else "export"
    if fmt == "sql_update":
        default = f"{default}_マスキング用UPDATE"
    file_name = safe_file_name(req.file_name or f"{default}_{now:%Y%m%d_%H%M%S}", ext)
    job_id = uuid.uuid4().hex
    job = ExportJob(
        job_id=job_id,
        request=req,
        file_name=file_name,
        path=export_dir() / f"{job_id}{ext}",
        tables=[TableProgress(schema_name=t.schema_name, table=t.table) for t in req.tables],
    )
    with _jobs_lock:
        _jobs[job_id] = job
    threading.Thread(target=_run_job, args=(job,), name=f"export-{job_id[:8]}", daemon=True).start()
    return job


def _run_job(job: ExportJob) -> None:
    job.status = "running"
    job.started_at = dt.datetime.now()
    req = job.request
    opts = req.options
    tmp_path = job.path.with_name(job.path.name + ".part")
    info_rows: list[list[Any]] = []
    try:
        if opts.format != "xlsx":
            _run_sql_job(job, tmp_path)
            tmp_path.replace(job.path)
            job.status = "done"
            return
        book = ExcelBook(opts)
        with store.open_session(req.connection_id) as (dialect, conn, profile):
            for spec, progress in zip(req.tables, job.tables):
                progress.status = "running"
                meta = dialect.list_columns(conn, spec.schema_name, spec.table)
                columns = prepare_columns(spec, meta, opts)
                if opts.count_rows:
                    with conn.cursor() as cur:
                        cur.execute(dialect.build_count(spec))
                        total = int(cur.fetchone()[0])
                    progress.total_rows = min(total, spec.max_rows) if spec.max_rows else total
                written = _export_table(job, book, dialect, conn, spec, columns, progress)
                progress.status = "done"
                info_rows.append(_info_for_table(spec, columns, written))
        if opts.include_info_sheet:
            book.add_info_sheet(_info_sheet_rows(job, profile, info_rows))
        book.save(tmp_path)
        tmp_path.replace(job.path)
        job.status = "done"
        job.message = f"{sum(t.rows_written for t in job.tables):,} 件を出力しました。"
    except Cancelled:
        job.status = "cancelled"
        job.message = "エクスポートを中止しました。"
    except DbError as exc:
        job.status = "error"
        job.message = str(exc)
    except Exception as exc:  # noqa: BLE001
        log.exception("エクスポート処理で予期しないエラーが発生しました")
        job.status = "error"
        job.message = f"予期しないエラーが発生しました: {exc}"
    finally:
        for t in job.tables:
            if t.status == "running":
                t.status = "error"
        job.finished_at = dt.datetime.now()
        if tmp_path.exists():
            try:
                tmp_path.unlink()
            except OSError:
                pass


def _export_table(
    job: ExportJob,
    book: ExcelBook,
    dialect: Any,
    conn: Any,
    spec: TableExportSpec,
    columns: list[PreparedColumn],
    progress: TableProgress,
) -> int:
    header_count = 2 if book.options.header_mode == "both" else 1
    per_sheet = EXCEL_MAX_ROWS - header_count
    title = spec.sheet_name or spec.table
    cur = dialect.open_stream(conn, dialect.build_select(spec))
    try:
        batch = [transform_row(r, columns) for r in cur.fetchmany(FETCH_SIZE)]
        sheet_no = 1
        ws = book.new_sheet(title, columns, batch)
        in_sheet = 0
        while batch:
            for values in batch:
                if in_sheet >= per_sheet:
                    # シートの行数上限に達したら次のシートへ
                    book.finish_sheet(ws, len(columns), in_sheet)
                    sheet_no += 1
                    ws = book.new_sheet(f"{title}_{sheet_no}", columns, batch)
                    in_sheet = 0
                book.append(ws, values)
                in_sheet += 1
            progress.rows_written += len(batch)
            if job.cancel_event.is_set():
                raise Cancelled()
            batch = [transform_row(r, columns) for r in cur.fetchmany(FETCH_SIZE)]
        book.finish_sheet(ws, len(columns), in_sheet)
    finally:
        cur.close()
    return progress.rows_written


def _info_for_table(spec: TableExportSpec, columns: list[PreparedColumn], written: int) -> list[Any]:
    masked = [f"{c.name}（{c.mask_label}）" for c in columns if c.mask_label]
    return [
        f"{spec.schema_name}.{spec.table}",
        written,
        spec.where or "",
        spec.order_by or "",
        ", ".join(c.name for c in columns),
        "、".join(masked) if masked else "なし",
    ]


def _info_sheet_rows(job: ExportJob, profile: Any, tables: list[list[Any]]) -> list[list[Any]]:
    if profile.db_type == "oracle":
        target = profile.dsn if profile.oracle_connect_type == "dsn" else f"{profile.host}:{profile.port or 1521}/{profile.database}"
    else:
        target = f"{profile.host}:{profile.port or 5432}/{profile.database}"
    rows: list[list[Any]] = [
        ["エクスポート情報"],
        [],
        ["出力日時", job.started_at.strftime("%Y/%m/%d %H:%M:%S") if job.started_at else ""],
        ["接続名", profile.name],
        ["DB種別", "Oracle" if profile.db_type == "oracle" else "PostgreSQL"],
        ["接続先", target],
        ["接続ユーザー", profile.username],
        ["ファイル名", job.file_name],
        [],
        ["対象表", "出力件数", "抽出条件(WHERE)", "並び順(ORDER BY)", "出力列", "マスキング"],
    ]
    rows.extend(tables)
    return rows


def cancel_job(job_id: str) -> ExportJob:
    job = get_job(job_id)
    job.cancel_event.set()
    return job


# ---------------------------------------------------------------------------
# SQL 出力（INSERT 文 / UPDATE 文）
# ---------------------------------------------------------------------------

INSERT_ROWS_PER_STATEMENT_PG = 100


def _target_label(profile: Any) -> str:
    if profile.db_type == "oracle":
        target = profile.dsn if profile.oracle_connect_type == "dsn" else f"{profile.host}:{profile.port or 1521}/{profile.database}"
        return f"Oracle {target}"
    return f"PostgreSQL {profile.host}:{profile.port or 5432}/{profile.database}"


def _table_ref(dialect: Any, spec: TableExportSpec, opts: ExportOptions) -> str:
    table = dialect.quote_ident(spec.table)
    if not opts.sql_qualify_schema:
        return table
    schema = opts.sql_target_schema.strip() or spec.schema_name
    return f"{dialect.quote_ident(schema)}.{table}"


def _comment(text: str) -> str:
    return "\n".join("-- " + line if line else "--" for line in text.splitlines())


def _plans(dialect: Any, meta: list[ColumnInfo], spec: TableExportSpec, opts: ExportOptions) -> list[sqlgen.ColumnPlan]:
    by_name = {c.name: c for c in meta}
    plans = []
    for col in spec.columns:
        info = by_name.get(col.name)
        if info is None:
            raise DbError(f"列 {col.name} は {spec.schema_name}.{spec.table} に存在しません。")
        if col.mask is not None and col.mask.type not in ("none", *(r["type"] for r in RULE_DEFINITIONS)):
            raise DbError(f"不明なマスキング規則です: {col.mask.type}")
        plans.append(sqlgen.plan_column(sqlgen.classify(dialect.db_type, info), col.mask, opts.hash_salt))
    return plans


def _file_header(job: ExportJob, profile: Any, title: str, extra: list[str]) -> str:
    lines = [
        "=" * 70,
        f"{title}（db-exp で生成）",
        f"生成日時 : {job.started_at:%Y/%m/%d %H:%M:%S}" if job.started_at else "",
        f"出力元   : {profile.name}（{_target_label(profile)} / ユーザー {profile.username}）",
        f"対象表   : {len(job.request.tables)} 表",
        *extra,
        "=" * 70,
    ]
    return _comment("\n".join(x for x in lines if x is not None))


def _session_setup(db_type: str, opts: ExportOptions) -> list[str]:
    if db_type == "oracle":
        return [
            "-- SQL*Plus / SQLcl / SQL Developer 用の設定（& を置換変数として扱わない・空行を許可・エラー時は中止）",
            "SET DEFINE OFF",
            "SET SQLBLANKLINES ON",
            "WHENEVER SQLERROR EXIT ROLLBACK",
            "",
        ]
    encoding = "SJIS" if opts.sql_encoding == "cp932" else "UTF8"
    return [
        "-- psql で実行する場合: psql -v ON_ERROR_STOP=1 -f <このファイル>",
        f"SET client_encoding = '{encoding}';",
        "SET standard_conforming_strings = on;",
        "BEGIN;",
        "",
    ]


def _run_sql_job(job: ExportJob, tmp_path: Path) -> None:
    req = job.request
    opts = req.options
    with store.open_session(req.connection_id) as (dialect, conn, profile):
        with open(tmp_path, "w", encoding=opts.sql_encoding, errors="replace", newline="\n") as out:
            if opts.format == "sql_insert":
                total = _write_inserts(job, out, dialect, conn, profile)
                job.message = f"{total:,} 件の INSERT 文を出力しました。"
            else:
                count = _write_updates(job, out, dialect, conn, profile)
                job.message = f"{count} 表の UPDATE 文を出力しました。"


def _mask_summary(plans: list[sqlgen.ColumnPlan]) -> str:
    masked = [f"{p.col.name}（{p.rule_label}）" for p in plans if p.rule is not None]
    return "、".join(masked) if masked else "なし"


def _write_inserts(job: ExportJob, out: Any, dialect: Any, conn: Any, profile: Any) -> int:
    req = job.request
    opts = req.options
    ora = dialect.db_type == "oracle"
    out.write(_file_header(job, profile, "マスク済みデータ（INSERT 文）", [
        f"COMMIT   : {opts.sql_commit_every:,} 件ごと" if opts.sql_commit_every else "COMMIT   : 最後に 1 回",
    ]) + "\n\n")
    setup = _session_setup(dialect.db_type, opts)
    if ora:
        # 1 行ごとの「1行が作成されました。」の表示を抑止する（大量件数で遅くなるため）
        setup.insert(-1, "SET FEEDBACK OFF")
    out.write("\n".join(setup) + "\n")
    grand_total = 0
    warnings: dict[str, int] = {}
    for spec, progress in zip(req.tables, job.tables):
        progress.status = "running"
        meta = dialect.list_columns(conn, spec.schema_name, spec.table)
        plans = _plans(dialect, meta, spec, opts)
        if opts.count_rows:
            with conn.cursor() as cur:
                cur.execute(dialect.build_count(spec))
                cnt = int(cur.fetchone()[0])
            progress.total_rows = min(cnt, spec.max_rows) if spec.max_rows else cnt
        ref = _table_ref(dialect, spec, opts)
        cols = ", ".join(dialect.quote_ident(p.col.name) for p in plans)
        notes = [f"{p.col.name}: {p.note}" for p in plans if p.note]
        out.write("\n" + _comment("\n".join(x for x in [
            "-" * 70,
            f"表       : {spec.schema_name}.{spec.table}",
            f"マスク   : {_mask_summary(plans)}",
            f"抽出条件 : {spec.where}" if spec.where else "",
            f"並び順   : {spec.order_by}" if spec.order_by else "",
            f"最大件数 : {spec.max_rows:,}" if spec.max_rows else "",
            *(f"注意     : {n}" for n in notes),
            "-" * 70,
        ] if x)) + "\n")
        if opts.sql_delete_before_insert:
            out.write(f"DELETE FROM {ref};\n")

        prefix = f"INSERT INTO {ref} ({cols}) VALUES"
        pending: list[str] = []
        since_commit = 0

        def flush() -> None:
            if not pending:
                return
            if ora:
                out.write("".join(f"{prefix} ({v});\n" for v in pending))
            else:
                out.write(prefix + "\n  " + ",\n  ".join(f"({v})" for v in pending) + ";\n")
            pending.clear()

        cur = dialect.open_stream(conn, dialect.build_select(spec))
        try:
            while True:
                rows = cur.fetchmany(FETCH_SIZE)
                if not rows:
                    break
                for row in rows:
                    values = []
                    for value, plan in zip(row, plans):
                        v = sqlgen.apply_plan(plan, value)
                        try:
                            values.append(sqlgen.literal(dialect.db_type, plan.col, v))
                        except sqlgen.LiteralError as exc:
                            key = f"{spec.table}.{plan.col.name}: {exc}"
                            warnings[key] = warnings.get(key, 0) + 1
                            values.append("NULL")
                    pending.append(", ".join(values))
                    since_commit += 1
                    if ora or len(pending) >= INSERT_ROWS_PER_STATEMENT_PG:
                        flush()
                    if opts.sql_commit_every and since_commit >= opts.sql_commit_every:
                        flush()
                        out.write("COMMIT;\n" if ora else "COMMIT;\nBEGIN;\n")
                        since_commit = 0
                progress.rows_written += len(rows)
                if job.cancel_event.is_set():
                    raise Cancelled()
            flush()
        finally:
            cur.close()
        grand_total += progress.rows_written
        progress.status = "done"
        out.write(f"-- {spec.schema_name}.{spec.table}: {progress.rows_written:,} 件\n")

    out.write("\nCOMMIT;\n")
    if warnings:
        out.write("\n" + _comment("警告:\n" + "\n".join(f"  {k}（{n:,} 件）" for k, n in warnings.items())) + "\n")
    return grand_total


def _write_updates(job: ExportJob, out: Any, dialect: Any, conn: Any, profile: Any) -> int:
    req = job.request
    opts = req.options
    db = dialect.db_type
    extra = [
        "",
        "【重要】本番環境では絶対に実行しないでください。",
        "        本番データをコピーしたテスト環境で、表の全行をその場でマスクします（元に戻せません）。",
        "・抽出条件（WHERE）・最大件数は適用されません（全行が対象）。",
        "・「出力する」のチェックを外した列は UPDATE されません（元の値のまま残ります）。",
        "・文中の個人情報のマスク（備考欄等）は簡易版です（メール・電話・郵便番号・10桁以上の数字・敬称の前の語のみ）。",
    ]
    if opts.hash_salt:
        extra.append("・ハッシュ化・ダミー氏名のソルトがこのファイルに含まれます。取り扱いに注意してください。")
    if db == "oracle":
        extra.append("・ハッシュ化・ダミー氏名には Oracle 12c 以降の STANDARD_HASH を使用します。")
    out.write(_file_header(job, profile, "マスキング用 UPDATE 文", extra) + "\n\n")
    out.write("\n".join(_session_setup(db, opts)) + "\n")
    count = 0
    for spec, progress in zip(req.tables, job.tables):
        progress.status = "running"
        meta = dialect.list_columns(conn, spec.schema_name, spec.table)
        plans = _plans(dialect, meta, spec, opts)
        included = {c.name for c in spec.columns}
        excluded = [c.name for c in meta if c.name not in included]
        sets = []
        for plan in plans:
            quoted = dialect.quote_ident(plan.col.name)
            expr = sqlgen.update_expr(db, plan, quoted, opts.hash_salt)
            if expr is not None:
                sets.append((quoted, expr))
        notes = [f"{p.col.name}: {p.note}" for p in plans if p.note]
        if any(p.col.is_pk and p.action != "keep" for p in plans):
            notes.append("主キーの列を更新します（外部キー制約がある場合はエラーになります）")
        out.write("\n" + _comment("\n".join(x for x in [
            "-" * 70,
            f"表       : {spec.schema_name}.{spec.table}",
            f"マスク   : {_mask_summary(plans)}",
            f"対象外の列（更新しない）: {', '.join(excluded)}" if excluded else "",
            *(f"注意     : {n}" for n in notes),
            "-" * 70,
        ] if x)) + "\n")
        if not sets:
            out.write("-- マスク対象の列がないため、この表の UPDATE 文はありません。\n")
        else:
            body = ",\n       ".join(f"{q} = {e}" for q, e in sets)
            out.write(sqlgen.wrap_sql(f"UPDATE {_table_ref(dialect, spec, opts)}\n   SET {body};") + "\n")
            count += 1
        progress.status = "done"
        if job.cancel_event.is_set():
            raise Cancelled()
    out.write("\nCOMMIT;\n")
    return count
