"""ドキュメント（Excel / CSV / テキスト / Markdown）のスキャンと、マスク済み Markdown の出力。

- 表になっている部分（CSV 全体、Excel の見出し行付きの範囲、Markdown の表）は列ごとに判定する。
- それ以外の文章（テキスト、Excel の自由記述、表の中で判定対象外の列のセル）は、
  文中の個人情報（textscan）を検出する。
- 出力は元のファイルごとの Markdown と「マスキング結果一覧.md」を ZIP にまとめる。
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
import shutil
import time
import unicodedata
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path, PurePosixPath
from typing import Any, Iterator

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .categories import CATEGORY_BY_CODE, check_email, check_phone
from .config import data_dir, load_settings
from .db import DbError
from .detector import analyze_column
from .findings import column_finding, effective_rule, make_id, masked_examples
from .masking import RULE_DEFINITIONS, build_masker
from .schemas import Decision, DocFileInfo, DocScanResult, Finding, MaskRule
from .textscan import TextMatch, find_pii

SUPPORTED = {
    ".xlsx": "excel",
    ".xlsm": "excel",
    ".csv": "csv",
    ".tsv": "csv",
    ".txt": "text",
    ".log": "text",
    ".md": "markdown",
    ".markdown": "markdown",
}
KIND_LABELS = {"excel": "Excel", "csv": "CSV", "text": "テキスト", "markdown": "Markdown"}
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_TOTAL_BYTES = 300 * 1024 * 1024
SAMPLE_ROWS = 500
REPORT_NAME = "マスキング結果一覧.md"
_RULE_LABELS = {r["type"]: r["label"] for r in RULE_DEFINITIONS}


# ---------------------------------------------------------------------------
# 読み込み
# ---------------------------------------------------------------------------


@dataclass
class TableBlock:
    section: str | None
    loc: str
    headers: list[str]
    col_labels: list[str]
    rows: list[list[str]]
    # Markdown の表の区切り行（配置指定を保つため）
    md_align: str | None = None


@dataclass
class LineBlock:
    section: str | None
    lines: list[tuple[str, str]]  # (場所, 本文)
    # True: Markdown の行としてそのまま出力する / False: プレーンテキストとしてエスケープする
    markdown: bool = False


@dataclass
class ParsedDoc:
    name: str
    kind: str
    size: int
    blocks: list[TableBlock | LineBlock] = field(default_factory=list)
    encoding: str | None = None


def read_text(path: Path) -> tuple[str, str]:
    """日本語でよく使われる文字コードを判定して読み込む。"""
    data = path.read_bytes()
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace"), "UTF-8（BOM付き）"
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace"), "UTF-16"
    for enc, label in (("utf-8", "UTF-8"), ("cp932", "Shift_JIS"), ("euc_jp", "EUC-JP")):
        try:
            return data.decode(enc), label
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "UTF-8（一部読めない文字あり）"


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        return value.strftime("%Y/%m/%d") if value.time() == dt.time() else value.strftime("%Y/%m/%d %H:%M:%S")
    if isinstance(value, dt.date):
        return value.strftime("%Y/%m/%d")
    if isinstance(value, dt.time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value).replace("\r\n", "\n").replace("\r", "\n")


_NUMERIC = re.compile(r"^[-+]?[\d,]+(\.\d+)?%?$")


def _looks_like_header(cells: list[str]) -> bool:
    """見出し行らしいか（数値・メールアドレス・電話番号など、データらしい値を含まない）。"""
    filled = [c.strip() for c in cells if c.strip()]
    if len(filled) < 2:
        return False
    for c in filled:
        n = unicodedata.normalize("NFKC", c)
        if len(n) > 40 or _NUMERIC.match(n) or check_email(n) or check_phone(n) >= 0.7:
            return False
        if re.match(r"^\d{4}[/\-]\d{1,2}[/\-]\d{1,2}", n):
            return False
    return len(set(filled)) == len(filled)


def parse_csv(path: Path, name: str) -> ParsedDoc:
    text, enc = read_text(path)
    doc = ParsedDoc(name, "csv", path.stat().st_size, encoding=enc)
    first = text.split("\n", 1)[0]
    delimiter = "\t" if path.suffix.lower() == ".tsv" or first.count("\t") > first.count(",") else ","
    rows = [list(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)]
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return doc
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    if _looks_like_header(rows[0]):
        headers, data, start = rows[0], rows[1:], 2
    else:
        headers, data, start = [f"列{i + 1}" for i in range(width)], rows, 1
    doc.blocks.append(
        TableBlock(None, f"{start}〜{start + len(data) - 1}行目", headers, [f"{i + 1}列目" for i in range(width)], data)
    )
    return doc


_MD_TABLE_ROW = re.compile(r"^\s*\|.*\|\s*$")
_MD_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")


def _split_md_row(line: str) -> list[str]:
    body = line.strip()
    if body.startswith("|"):
        body = body[1:]
    if body.endswith("|") and not body.endswith("\\|"):
        body = body[:-1]
    cells = re.split(r"(?<!\\)\|", body)
    return [c.strip().replace("\\|", "|").replace("<br>", "\n") for c in cells]


def parse_text(path: Path, name: str, markdown: bool) -> ParsedDoc:
    text, enc = read_text(path)
    doc = ParsedDoc(name, "markdown" if markdown else "text", path.stat().st_size, encoding=enc)
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    if lines and lines[-1] == "":
        lines.pop()
    current: list[tuple[str, str]] = []
    i = 0
    in_fence = False
    while i < len(lines):
        line = lines[i]
        if markdown and line.lstrip().startswith(("```", "~~~")):
            in_fence = not in_fence
        if (
            markdown
            and not in_fence
            and _MD_TABLE_ROW.match(line)
            and i + 1 < len(lines)
            and _MD_TABLE_SEP.match(lines[i + 1])
        ):
            if current:
                doc.blocks.append(LineBlock(None, current, markdown=True))
                current = []
            headers = _split_md_row(line)
            align = lines[i + 1]
            j = i + 2
            rows = []
            while j < len(lines) and _MD_TABLE_ROW.match(lines[j]):
                cells = _split_md_row(lines[j])
                rows.append((cells + [""] * len(headers))[: len(headers)])
                j += 1
            doc.blocks.append(
                TableBlock(None, f"{i + 1}〜{j}行目の表", headers, [f"{k + 1}列目" for k in range(len(headers))], rows, align)
            )
            i = j
            continue
        current.append((f"{i + 1}行目", line))
        i += 1
    if current:
        doc.blocks.append(LineBlock(None, current, markdown=markdown))
    return doc


def _excel_segments(rows: list[tuple[int, list[str]]]) -> Iterator[list[tuple[int, list[str]]]]:
    """空行で区切られた範囲ごとに分ける。"""
    seg: list[tuple[int, list[str]]] = []
    for r in rows:
        if any(c.strip() for c in r[1]):
            seg.append(r)
        elif seg:
            yield seg
            seg = []
    if seg:
        yield seg


def _excel_table_at(seg: list[tuple[int, list[str]]], k: int) -> tuple[int, int] | None:
    """seg[k] を見出し行とする表になっているか。なっていれば列の範囲を返す。"""
    header = seg[k][1]
    data = seg[k + 1 :]
    if not data or not _looks_like_header(header):
        return None
    filled = [i for i, c in enumerate(header) if c.strip()]
    lo, hi = min(filled), max(filled)
    inside = sum(1 for _, r in data if all(not c.strip() for i, c in enumerate(r) if i < lo or i > hi))
    multi = sum(1 for _, r in data if sum(1 for c in r[lo : hi + 1] if c.strip()) >= 2)
    if inside / len(data) >= 0.8 and multi >= max(1, len(data) * 0.3):
        return lo, hi
    return None


def parse_excel(path: Path, name: str) -> ParsedDoc:
    doc = ParsedDoc(name, "excel", path.stat().st_size)
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise DbError(f"Excel ファイルを開けません: {exc}") from exc
    try:
        for ws in wb.worksheets:
            rows: list[tuple[int, list[str]]] = []
            for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                cells = [cell_text(v) for v in row]
                while cells and not cells[-1].strip():
                    cells.pop()
                rows.append((idx, cells))
            width = max((len(r) for _, r in rows), default=0)
            rows = [(i, r + [""] * (width - len(r))) for i, r in rows]
            for seg in _excel_segments(rows):
                table = None
                for k in range(min(3, len(seg) - 1)):
                    span = _excel_table_at(seg, k)
                    if span:
                        table = (k, span)
                        break
                if table is None:
                    doc.blocks.append(LineBlock(ws.title, [_excel_line(ws.title, i, r) for i, r in seg]))
                    continue
                k, (lo, hi) = table
                if k:
                    doc.blocks.append(LineBlock(ws.title, [_excel_line(ws.title, i, r) for i, r in seg[:k]]))
                header_row = seg[k][0]
                last_row = seg[-1][0]
                doc.blocks.append(
                    TableBlock(
                        ws.title,
                        f"{get_column_letter(lo + 1)}{header_row}:{get_column_letter(hi + 1)}{last_row}",
                        [c.strip() for c in seg[k][1][lo : hi + 1]],
                        [f"{get_column_letter(c + 1)}列" for c in range(lo, hi + 1)],
                        [r[lo : hi + 1] for _, r in seg[k + 1 :]],
                    )
                )
    finally:
        wb.close()
    return doc


def _excel_line(sheet: str, row: int, cells: list[str]) -> tuple[str, str]:
    # セルの区切りは Markdown の表と誤認されない全角の縦線にする
    return f"{row}行目", " ｜ ".join(c.strip().replace("\n", " ") for c in cells if c.strip())


def parse_document(path: Path, name: str) -> ParsedDoc:
    kind = SUPPORTED.get(path.suffix.lower())
    if kind == "excel":
        return parse_excel(path, name)
    if kind == "csv":
        return parse_csv(path, name)
    if kind in ("text", "markdown"):
        return parse_text(path, name, kind == "markdown")
    raise DbError("未対応のファイル形式です（対応: .xlsx .csv .tsv .txt .md）。")


# ---------------------------------------------------------------------------
# 判定
# ---------------------------------------------------------------------------


def _text_id(file: str, m: TextMatch, fragment: str) -> str:
    """文中の検出結果をまとめる単位の ID。

    自動判定はファイル×種類ごとに 1 件、要確認は同じ語句ごとに 1 件にまとめる。
    """
    if m.level == "auto":
        return make_id("doc", "text", file, m.category, "auto")
    key = re.sub(r"\s+", "", unicodedata.normalize("NFKC", fragment))
    return make_id("doc", "text", file, m.category, "review", key)


def _skip_cell(text: str) -> bool:
    t = text.strip()
    return len(t) < 2 or (t.isdigit() and len(t) <= 6)


@dataclass
class Analysis:
    columns: dict[tuple[int, int], Finding] = field(default_factory=dict)
    texts: dict[str, Finding] = field(default_factory=dict)

    def findings(self) -> list[Finding]:
        return [*self.columns.values(), *self.texts.values()]


def _block_place(block: TableBlock | LineBlock) -> str:
    return f"シート「{block.section}」" if block.section else ""


def _add_text(an: Analysis, doc: ParsedDoc, place: str, text: str) -> None:
    for m in find_pii(text):
        fragment = text[m.start : m.end]
        fid = _text_id(doc.name, m, fragment)
        cat = CATEGORY_BY_CODE[m.category]
        f = an.texts.get(fid)
        if f is None:
            rule = MaskRule(type=cat.rule, params=dict(cat.params))
            f = an.texts[fid] = Finding(
                id=fid, source="doc", target="text", level=m.level, category=cat.code, category_label=cat.label,
                score=m.score, reasons=[m.reason], rule=rule, file=doc.name, location=place,
                text=fragment if m.level == "review" else None, occurrences=0,
            )
        f.occurrences += 1
        f.score = max(f.score, m.score)
        if m.reason not in f.reasons and len(f.reasons) < 4:
            f.reasons.append(m.reason)
        if len(f.samples) < 3 and fragment not in f.samples:
            f.samples.append(fragment)
            f.masked_samples = masked_examples(f.rule, f.samples)


def analyze_document(doc: ParsedDoc) -> Analysis:
    an = Analysis()
    for bi, block in enumerate(doc.blocks):
        place = _block_place(block)
        if isinstance(block, TableBlock):
            for ci, header in enumerate(block.headers):
                values = [r[ci] for r in block.rows[:SAMPLE_ROWS]]
                verdict = analyze_column(header, None, "", values, context="doc")
                if verdict.level is None:
                    continue
                loc = " ".join(x for x in (place, block.loc, block.col_labels[ci]) if x)
                an.columns[(bi, ci)] = column_finding(
                    source="doc", file=doc.name, column=header or block.col_labels[ci], verdict=verdict,
                    samples=values, location=loc, id_parts=(bi, ci),
                )
            for ri, row in enumerate(block.rows):
                for ci, value in enumerate(row):
                    if (bi, ci) not in an.columns and not _skip_cell(value):
                        _add_text(an, doc, " ".join(x for x in (place, block.loc) if x), value)
        else:
            for loc, text in block.lines:
                if not _skip_cell(text):
                    _add_text(an, doc, " ".join(x for x in (place, loc) if x), text)
    # 文中の検出結果の場所は「最初に見つかった場所（ほか N 件）」にする
    for f in an.texts.values():
        if f.occurrences > 1:
            f.location = f"{f.location} ほか {f.occurrences - 1} 件"
    return an


# ---------------------------------------------------------------------------
# アップロードの保存
# ---------------------------------------------------------------------------


def uploads_dir() -> Path:
    path = data_dir() / "uploads"
    path.mkdir(parents=True, exist_ok=True)
    return path


def safe_relative_name(name: str) -> str:
    """フォルダ付きのファイル名から、危険な要素（..・絶対パス・制御文字）を取り除く。"""
    parts = [p for p in re.split(r"[\\/]+", name or "") if p and p not in (".", "..")]
    parts = [re.sub(r'[\x00-\x1f:*?"<>|]', "_", p).strip(" .") or "_" for p in parts]
    return "/".join(parts) or "file"


def cleanup_uploads() -> None:
    retention = int(load_settings().get("export_retention_seconds") or 0)
    if retention <= 0:
        return
    limit = time.time() - retention
    for p in uploads_dir().iterdir():
        try:
            if p.is_dir() and p.stat().st_mtime < limit:
                shutil.rmtree(p, ignore_errors=True)
        except OSError:
            pass


@dataclass
class StoredFile:
    name: str
    path: Path


def _scan_dir(scan_id: str) -> Path:
    if not re.fullmatch(r"[0-9a-f]{32}", scan_id):
        raise DbError("スキャン結果が見つかりません。")
    return uploads_dir() / scan_id


def store_files(files: list[tuple[str, bytes]]) -> tuple[str, list[StoredFile]]:
    cleanup_uploads()
    total = sum(len(data) for _, data in files)
    if total > MAX_TOTAL_BYTES:
        raise DbError(f"ファイルの合計サイズが大きすぎます（上限 {MAX_TOTAL_BYTES // 1024 // 1024}MB）。")
    scan_id = uuid.uuid4().hex
    folder = _scan_dir(scan_id)
    folder.mkdir(parents=True)
    stored, used = [], set()
    for i, (raw_name, data) in enumerate(files):
        name = safe_relative_name(raw_name)
        base, n = name, 2
        while name in used:
            stem, dot, ext = base.rpartition(".")
            name = f"{stem}({n}).{ext}" if dot else f"{base}({n})"
            n += 1
        used.add(name)
        path = folder / f"{i:04d}{Path(name).suffix.lower()}"
        path.write_bytes(data)
        stored.append(StoredFile(name, path))
    (folder / "meta.json").write_text(
        json.dumps([{"name": s.name, "file": s.path.name} for s in stored], ensure_ascii=False), encoding="utf-8"
    )
    return scan_id, stored


def load_stored(scan_id: str) -> list[StoredFile]:
    folder = _scan_dir(scan_id)
    meta = folder / "meta.json"
    if not meta.exists():
        raise DbError("スキャン結果が見つかりません（保持期間を過ぎた可能性があります）。もう一度ファイルを読み込んでください。")
    return [StoredFile(m["name"], folder / m["file"]) for m in json.loads(meta.read_text(encoding="utf-8"))]


def delete_scan(scan_id: str) -> None:
    shutil.rmtree(_scan_dir(scan_id), ignore_errors=True)


def _parse_stored(f: StoredFile) -> ParsedDoc:
    if f.path.suffix.lower() not in SUPPORTED:
        raise DbError("未対応のファイル形式です（対応: .xlsx .csv .tsv .txt .md）。")
    size = f.path.stat().st_size
    if size > MAX_FILE_BYTES:
        raise DbError(f"ファイルが大きすぎます（上限 {MAX_FILE_BYTES // 1024 // 1024}MB）。大量データは DB から直接出力してください。")
    return parse_document(f.path, f.name)


def scan_files(files: list[tuple[str, bytes]]) -> DocScanResult:
    scan_id, stored = store_files(files)
    infos: list[DocFileInfo] = []
    findings: list[Finding] = []
    for f in stored:
        kind = SUPPORTED.get(f.path.suffix.lower(), "unknown")
        info = DocFileInfo(name=f.name, kind=kind, size=f.path.stat().st_size)
        try:
            doc = _parse_stored(f)
            info.blocks = len(doc.blocks)
            info.tables = sum(1 for b in doc.blocks if isinstance(b, TableBlock))
            info.lines = sum(len(b.rows) if isinstance(b, TableBlock) else len(b.lines) for b in doc.blocks)
            findings.extend(analyze_document(doc).findings())
        except DbError as exc:
            info.error = str(exc)
        except Exception as exc:  # noqa: BLE001
            info.error = f"読み込めませんでした: {exc}"
        infos.append(info)
    return DocScanResult(scan_id=scan_id, files=infos, findings=findings)


# ---------------------------------------------------------------------------
# Markdown 出力
# ---------------------------------------------------------------------------


class Masker:
    """人の判断を反映してセル・行をマスクする（件数も数える）。"""

    def __init__(self, decisions: dict[str, Decision], salt: str):
        self.decisions = decisions
        self.salt = salt
        self.counts: dict[str, int] = {}
        self._cache: dict[str, Any] = {}

    def _fn(self, rule: MaskRule | None) -> Any:
        if rule is None:
            return None
        key = rule.model_dump_json()
        if key not in self._cache:
            self._cache[key] = build_masker(rule, self.salt)
        return self._cache[key]

    def column(self, finding: Finding, value: str) -> str:
        rule = effective_rule(finding, self.decisions)
        fn = self._fn(rule)
        if fn is None or not value.strip():
            return value
        self.counts[finding.id] = self.counts.get(finding.id, 0) + 1
        out = fn(value)
        return "" if out is None else str(out)

    def text(self, an: Analysis, file: str, text: str) -> str:
        if _skip_cell(text):
            return text
        out, pos = [], 0
        for m in find_pii(text):
            fragment = text[m.start : m.end]
            fid = _text_id(file, m, fragment)
            finding = an.texts.get(fid)
            rule = effective_rule(finding, self.decisions) if finding else None
            fn = self._fn(rule)
            if fn is None:
                continue
            masked = fn(fragment)
            out.append(text[pos : m.start])
            out.append("" if masked is None else str(masked))
            pos = m.end
            self.counts[fid] = self.counts.get(fid, 0) + 1
        out.append(text[pos:])
        return "".join(out)


def _md_cell(text: str) -> str:
    return text.strip().replace("\\", "\\\\").replace("|", "\\|").replace("\n", "<br>")


_MD_LEADING = re.compile(r"^(\s*)([#>*+\-=|`~]|\d+[.)])")


def _md_plain_line(text: str) -> str:
    """プレーンテキストの 1 行を、Markdown として表示が崩れないようにする。"""
    text = text.replace("<", "&lt;")
    text = _MD_LEADING.sub(lambda m: m.group(1) + "\\" + m.group(2), text)
    return text


def render_markdown(doc: ParsedDoc, an: Analysis, masker: Masker, generated: str) -> str:
    out: list[str] = []
    if doc.kind == "markdown":
        out.append(f"<!-- db-exp でマスキング済み（{generated}）元ファイル: {doc.name} -->")
        out.append("")
    else:
        out.append(f"# {_md_plain_line(doc.name)}")
        out.append("")
        out.append(f"> db-exp でマスキング済み（{generated}）。元ファイル: {doc.name}（{KIND_LABELS.get(doc.kind, doc.kind)}）")
        out.append("")
    section = None
    for bi, block in enumerate(doc.blocks):
        if doc.kind == "excel" and block.section != section:
            section = block.section
            out.append(f"## {_md_plain_line(section or '')}")
            out.append("")
        if isinstance(block, TableBlock):
            headers = [h or label for h, label in zip(block.headers, block.col_labels)]
            out.append("| " + " | ".join(_md_cell(h) for h in headers) + " |")
            out.append(block.md_align or "| " + " | ".join("---" for _ in headers) + " |")
            for row in block.rows:
                cells = []
                for ci, value in enumerate(row):
                    finding = an.columns.get((bi, ci))
                    v = masker.column(finding, value) if finding else masker.text(an, doc.name, value)
                    cells.append(_md_cell(v))
                out.append("| " + " | ".join(cells) + " |")
            out.append("")
        else:
            for _, text in block.lines:
                masked = masker.text(an, doc.name, text)
                if block.markdown:
                    out.append(masked)
                elif not masked.strip():
                    out.append("")
                else:
                    # 改行を保つため、行末に半角空白 2 つ（Markdown の改行）を付ける
                    out.append(_md_plain_line(masked) + "  ")
            if not block.markdown:
                out.append("")
    return "\n".join(out).rstrip() + "\n"


def _output_name(name: str, used: set[str]) -> str:
    p = PurePosixPath(name)
    candidate = str(p.with_suffix(".md")) if p.suffix else f"{name}.md"
    if candidate in used or candidate == REPORT_NAME:
        candidate = str(p.with_name(f"{p.stem}_{p.suffix.lstrip('.')}.md"))
    n = 2
    base = candidate
    while candidate in used:
        candidate = base[:-3] + f"({n}).md"
        n += 1
    used.add(candidate)
    return candidate


def _decision_label(f: Finding, decisions: dict[str, Decision]) -> tuple[str, str]:
    """（判断、処理）の表示用文字列。"""
    d = decisions.get(f.id)
    rule = effective_rule(f, decisions)
    judged = "自動" if f.level == "auto" else ("人が判断" if d is not None else "未判断（安全のためマスク）")
    action = _RULE_LABELS.get(rule.type, rule.type) if rule else "マスクしない"
    return judged, action


def build_report_md(
    title: str, generated: str, findings: list[Finding], decisions: dict[str, Decision],
    counts: dict[str, int] | None = None, files: list[DocFileInfo] | None = None,
) -> str:
    auto = [f for f in findings if f.level == "auto"]
    review = [f for f in findings if f.level == "review"]
    masked_review = sum(1 for f in review if effective_rule(f, decisions) is not None)
    out = [
        f"# {title}",
        "",
        f"- 出力日時: {generated}",
        f"- マスキング対象（自動判定）: {len(auto)} 件",
        f"- 要確認（人が判断）: {len(review)} 件（マスクする {masked_review} 件 / マスクしない {len(review) - masked_review} 件）",
        "",
    ]
    if files:
        out += ["## 対象ファイル", "", "| ファイル | 形式 | 状態 |", "| --- | --- | --- |"]
        for info in files:
            state = info.error or f"表 {info.tables} / {info.lines:,} 行"
            out.append(f"| {_md_cell(info.name)} | {KIND_LABELS.get(info.kind, info.kind)} | {_md_cell(state)} |")
        out.append("")

    def table(rows: list[Finding]) -> list[str]:
        lines = [
            "| No | ファイル・表 | 場所 | 対象 | 種類 | 判断 | 処理 | 件数 | マスク後の例 | 判定理由 |",
            "| --: | --- | --- | --- | --- | --- | --- | --: | --- | --- |",
        ]
        for i, f in enumerate(rows, start=1):
            judged, action = _decision_label(f, decisions)
            target = f"列「{f.column}」" if f.target == "column" else "文中"
            if f.column_comment:
                target += f"（{f.column_comment}）"
            n = counts.get(f.id, 0) if counts is not None else f.occurrences
            example = " / ".join(masked_examples(effective_rule(f, decisions), f.samples)) if f.samples else ""
            lines.append(
                f"| {i} | {_md_cell(f.file)} | {_md_cell(f.location or '')} | {_md_cell(target)} | {_md_cell(f.category_label)} "
                f"| {judged} | {_md_cell(action)} | {n:,} | {_md_cell(example)} | {_md_cell('、'.join(f.reasons))} |"
            )
        return lines

    out += ["## マスキング対象（自動判定）", ""]
    out += table(auto) if auto else ["なし"]
    out += ["", "## 要確認（人が判断）", ""]
    out += table(review) if review else ["なし"]
    out.append("")
    return "\n".join(out)


def export_zip(scan_id: str, decisions: dict[str, Decision], salt: str = "") -> Path:
    """マスク済み Markdown と結果一覧を ZIP にまとめ、そのパスを返す。"""
    stored = load_stored(scan_id)
    generated = dt.datetime.now().strftime("%Y/%m/%d %H:%M:%S")
    masker = Masker(decisions, salt)
    all_findings: list[Finding] = []
    infos: list[DocFileInfo] = []
    out_path = _scan_dir(scan_id) / "output.zip"
    used: set[str] = set()
    with zipfile.ZipFile(out_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for f in stored:
            info = DocFileInfo(name=f.name, kind=SUPPORTED.get(f.path.suffix.lower(), "unknown"), size=f.path.stat().st_size)
            try:
                doc = _parse_stored(f)
            except DbError as exc:
                info.error = str(exc)
                infos.append(info)
                continue
            info.tables = sum(1 for b in doc.blocks if isinstance(b, TableBlock))
            info.lines = sum(len(b.rows) if isinstance(b, TableBlock) else len(b.lines) for b in doc.blocks)
            infos.append(info)
            an = analyze_document(doc)
            all_findings.extend(an.findings())
            zf.writestr(_output_name(f.name, used), render_markdown(doc, an, masker, generated))
        zf.writestr(REPORT_NAME, build_report_md("マスキング結果一覧", generated, all_findings, decisions, masker.counts, infos))
    return out_path
